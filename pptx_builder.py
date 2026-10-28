import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# Each theme: background, card, title text, body text, accent, font
THEMES = {
    "ocean":    {"name": "Ocean",    "bg": "0B1F3A", "card": "12305A", "title": "FFFFFF", "text": "DCE8F5", "accent": "38BDF8", "font": "Calibri"},
    "sunset":   {"name": "Sunset",   "bg": "2B0F2E", "card": "441A48", "title": "FFFFFF", "text": "F3DCEB", "accent": "FB923C", "font": "Calibri"},
    "forest":   {"name": "Forest",   "bg": "0F2A1D", "card": "17402C", "title": "FFFFFF", "text": "DDF1E5", "accent": "4ADE80", "font": "Calibri"},
    "midnight": {"name": "Midnight", "bg": "0A0A0F", "card": "1A1A24", "title": "FFFFFF", "text": "C9C9D6", "accent": "A78BFA", "font": "Calibri"},
    "paper":    {"name": "Paper",    "bg": "FFFFFF", "card": "EEF2F8", "title": "0F172A", "text": "334155", "accent": "2F5BFF", "font": "Calibri"},
}
W, H = 13.333, 7.5


def rgb(hex_code):
    return RGBColor.from_string(hex_code)


def lst(value):
    if not isinstance(value, list):
        return []
    return [str(x).strip() for x in value if str(x).strip()]


def rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(color)
    s.line.fill.background()
    return s


def text(slide, x, y, w, h, value, size, color, font, bold=False, italic=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = str(value or "")
    f = run.font
    f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, font
    f.color.rgb = rgb(color)
    return box


def bullet_list(slide, x, y, w, h, items, t, size):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(14)
        dot = p.add_run()
        dot.text = "●   "
        dot.font.size, dot.font.name = Pt(size - 6), t["font"]
        dot.font.color.rgb = rgb(t["accent"])
        run = p.add_run()
        run.text = item
        run.font.size, run.font.name = Pt(size), t["font"]
        run.font.color.rgb = rgb(t["text"])


def fit(items, base):
    total = sum(len(i) for i in items)
    if len(items) > 5 or total > 420:
        return base - 4
    if total > 300:
        return base - 2
    return base


def heading(slide, title, t):
    size = 34 if len(title) < 45 else 28
    text(slide, 0.8, 0.5, 11.7, 1.0, title, size, t["title"], t["font"], bold=True, anchor=MSO_ANCHOR.MIDDLE)
    rect(slide, 0.8, 1.55, 1.2, 0.07, t["accent"])


def layout_title(slide, s, t, closing=False):
    rect(slide, 0, 0, 0.35, H, t["accent"])
    rect(slide, 9.8, -1.6, 5.5, 5.5, t["card"], MSO_SHAPE.OVAL)
    rect(slide, 11.2, 5.2, 3.4, 3.4, t["card"], MSO_SHAPE.OVAL)
    title = s.get("title", "")
    size = (54 if closing else 48) if len(title) < 35 else 36
    text(slide, 1.2, 2.0, 9.5, 2.2, title, size, t["title"], t["font"], bold=True, anchor=MSO_ANCHOR.BOTTOM)
    rect(slide, 1.2, 4.4, 1.6, 0.08, t["accent"])
    text(slide, 1.2, 4.65, 9.5, 1.4, s.get("subtitle", ""), 22, t["text"], t["font"])


def layout_bullets(slide, s, t):
    heading(slide, s.get("title", ""), t)
    items = lst(s.get("bullets"))
    bullet_list(slide, 0.8, 2.0, 11.7, 4.8, items, t, fit(items, 24))


def layout_two(slide, s, t):
    heading(slide, s.get("title", ""), t)
    for i, (head, items) in enumerate([(s.get("left_title", ""), lst(s.get("left"))),
                                       (s.get("right_title", ""), lst(s.get("right")))]):
        x = 0.8 + i * 6.03
        rect(slide, x, 2.0, 5.7, 4.7, t["card"], MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.05
        text(slide, x + 0.3, 2.2, 5.1, 0.7, head, 22, t["accent"], t["font"], bold=True)
        bullet_list(slide, x + 0.3, 3.0, 5.1, 3.5, items, t, fit(items, 20) - 2)


def layout_stat(slide, s, t):
    heading(slide, s.get("title", ""), t)
    st = s.get("stat") or {}
    text(slide, 0.8, 2.1, 11.7, 2.3, st.get("value", ""), 96, t["accent"], t["font"], bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, 1.5, 4.6, 10.3, 1.5, st.get("label", ""), 26, t["text"], t["font"], align=PP_ALIGN.CENTER)


def layout_quote(slide, s, t):
    q = s.get("quote") or {}
    text(slide, 0.9, 0.5, 2, 2, "\u201C", 140, t["accent"], t["font"], bold=True)
    text(slide, 1.6, 2.2, 10.2, 3.0, q.get("text", ""), 30, t["title"], t["font"], italic=True, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, 1.6, 5.4, 10.2, 0.7, "\u2014 " + str(q.get("author", "")), 20, t["accent"], t["font"], bold=True)


LAYOUTS = {"bullets": layout_bullets, "two-column": layout_two, "stat": layout_stat, "quote": layout_quote}


def build_pptx(deck, theme_key="ocean"):
    t = THEMES.get(theme_key, THEMES["ocean"])
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]

    for i, s in enumerate(deck.get("slides", [])):
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(t["bg"])
        kind = s.get("layout", "bullets")
        if kind in ("title", "closing"):
            layout_title(slide, s, t, closing=(kind == "closing"))
        else:
            LAYOUTS.get(kind, layout_bullets)(slide, s, t)
            text(slide, 11.8, 6.95, 1.0, 0.4, str(i + 1), 12, t["text"], t["font"], align=PP_ALIGN.RIGHT)
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(s["notes"])

    out = io.BytesIO()
    prs.save(out)
    out.seek(0)
    return out
